"""Render the crypto intraday system design into the repository's DOCX template.

The Markdown file is the canonical, reviewable source. This renderer intentionally
supports only the Markdown constructs used by that document so the generated Word
artifact remains deterministic and easy to reproduce.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "crypto-intraday-system-design.md"
OUTPUT = ROOT / "artifacts" / "crypto-intraday-system-design.docx"
ARCHITECTURE_IMAGE = ROOT / "artifacts" / "diagrams" / "crypto-intraday-architecture.png"
SEQUENCE_IMAGE = ROOT / "artifacts" / "diagrams" / "crypto-intraday-decision-sequence.png"
TEMPLATE = Path(
    "/home/julius/.codex/plugins/cache/openai-curated-remote/"
    "openai-templates/0.1.1/skills/artifact-template-system-design/"
    "assets/reference.docx"
)

NAVY = "0E3152"
INK = RGBColor(23, 32, 51)
MUTED = RGBColor(88, 105, 128)
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F4F6F8"
WHITE = RGBColor(255, 255, 255)


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F6FA3")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, color: RGBColor | None = None, size: float | None = None) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            if color:
                run.font.color.rgb = color
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(size or 8.5)
            run.font.color.rgb = color or INK
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if color:
                run.font.color.rgb = color
            if size:
                run.font.size = Pt(size)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if color:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)


def normalize_markdown_cell(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.strip()


def add_styled_table(document: Document, rows: list[list[str]], *, font_size: float = 8.0):
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(normalize_markdown_cell(value))
            run.font.size = Pt(font_size)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                run.bold = True
                run.font.color.rgb = WHITE
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)
                run.font.color.rgb = INK
            else:
                run.font.color.rgb = INK
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_cover(document: Document) -> None:
    band = document.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = band.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=260, start=260, bottom=260, end=260)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("SYSTEM DESIGN RFC")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = WHITE

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(20)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(3)
    add_inline(title, "Crypto Intraday Trading System", color=INK)

    subtitle = document.add_paragraph(style="Title")
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Jev + Multi-Agent LLM")
    run.font.color.rgb = RGBColor(47, 111, 163)

    strap = document.add_paragraph()
    strap.paragraph_format.space_after = Pt(18)
    add_inline(
        strap,
        "Paper-only architecture for BTCUSDT perpetual futures · isolated 3x · deterministic risk control",
        color=MUTED,
        size=10,
    )

    metadata = document.add_table(rows=1, cols=3)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, (label, value) in enumerate(
        (
            ("STATUS", "M1 foundation"),
            ("OWNER", "Julius"),
            ("LAST UPDATED", "September 22, 2026"),
        )
    ):
        cell = metadata.cell(0, index)
        set_cell_shading(cell, PALE_BLUE)
        set_cell_margins(cell, top=140, start=150, bottom=140, end=150)
        paragraph = cell.paragraphs[0]
        label_run = paragraph.add_run(label + "\n")
        label_run.bold = True
        label_run.font.size = Pt(7.5)
        label_run.font.color.rgb = MUTED
        value_run = paragraph.add_run(value)
        value_run.bold = True
        value_run.font.size = Pt(10)
        value_run.font.color.rgb = INK

    document.add_paragraph().paragraph_format.space_after = Pt(14)
    details = document.add_table(rows=4, cols=2)
    details.alignment = WD_TABLE_ALIGNMENT.CENTER
    detail_rows = (
        ("Authors", "Julius; OpenAI Codex"),
        ("Reviewers", "Risk, engineering, and operations reviewers before implementation"),
        ("Related docs", "docs/architecture.md · docs/evaluation.md · docs/btc-paper-runbook.md"),
        ("Scope", "Standalone BTCUSDT perpetual paper system; the daily Donchian bot remains unchanged."),
    )
    for row, (label, value) in zip(details.rows, detail_rows):
        prevent_row_split(row)
        row.cells[0].width = Inches(1.3)
        for cell in row.cells:
            set_cell_margins(cell, top=100, start=130, bottom=100, end=130)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        label_run.bold = True
        label_run.font.size = Pt(8)
        label_run.font.color.rgb = MUTED
        value_run = row.cells[1].paragraphs[0].add_run(value)
        value_run.font.size = Pt(8.5)
        value_run.font.color.rgb = INK

    document.add_paragraph().paragraph_format.space_after = Pt(20)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(note, "Design artifact · not an authorization for live trading", color=MUTED, size=8)
    note.add_run().add_break(WD_BREAK.PAGE)


def configure_document(document: Document) -> None:
    properties = document.core_properties
    properties.title = "Crypto Intraday Trading System — Jev + Multi-Agent LLM"
    properties.subject = "Proposed architecture for a paper-only BTCUSDT perpetual trading system"
    properties.author = "Julius; OpenAI Codex"
    properties.keywords = "system design, paper trading, Jev, LLM, BTCUSDT, perpetual futures"
    properties.comments = "Generated from docs/crypto-intraday-system-design.md"

    normal = document.styles["normal"]
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.color.rgb = RGBColor(14, 49, 82)
        style.paragraph_format.keep_with_next = True

    for section in document.sections:
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        footer = section.footer
        for paragraph in footer.paragraphs:
            for run in paragraph.runs:
                if "Organization Name" in run.text or "System Design RFC" in run.text:
                    run.text = "Crypto Intraday Trading System | Proposed System Design"
                    run.font.color.rgb = MUTED


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        values = [value.strip() for value in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def is_special_line(line: str) -> bool:
    stripped = line.strip()
    return bool(
        not stripped
        or stripped.startswith("#")
        or stripped.startswith("```")
        or stripped.startswith("|")
        or stripped.startswith("- ")
        or stripped.startswith(">")
        or re.match(r"\d+\.\s", stripped)
    )


def add_code_block(document: Document, code: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), PALE_GRAY)
    properties.append(shading)
    run = paragraph.add_run("\n".join(code))
    run.font.name = "Liberation Mono"
    run.font.size = Pt(7.5)
    run.font.color.rgb = INK


def add_figure(document: Document, path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(path), width=Inches(7.0))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    run = caption_paragraph.add_run(caption)
    run.italic = True
    run.font.size = Pt(7.5)
    run.font.color.rgb = MUTED


def render_body(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    start = next(index for index, line in enumerate(lines) if line.startswith("## 1."))
    lines = lines[start:]
    index = 0
    mermaid_index = 0

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```mermaid"):
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                index += 1
            index += 1
            if mermaid_index == 0:
                add_figure(document, ARCHITECTURE_IMAGE, "Figure 1. Proposed system architecture and trust boundaries.")
            else:
                add_figure(document, SEQUENCE_IMAGE, "Figure 2. Five-second decision lifecycle and fail-closed paths.")
            mermaid_index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            add_code_block(document, code)
            continue

        if stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, stripped[3:], color=RGBColor(14, 49, 82))
            index += 1
            continue

        if stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_inline(paragraph, stripped[4:], color=RGBColor(14, 49, 82))
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_styled_table(document, rows, font_size=7.5 if len(rows[0]) >= 4 else 8.0)
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.first_line_indent = Inches(-0.13)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, "• " + stripped[2:])
            index += 1
            while index < len(lines) and lines[index].startswith("  ") and lines[index].strip():
                add_inline(paragraph, " " + lines[index].strip())
                index += 1
            continue

        if re.match(r"\d+\.\s", stripped):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, stripped)
            index += 1
            while index < len(lines) and lines[index].startswith("   ") and lines[index].strip():
                add_inline(paragraph, " " + lines[index].strip())
                index += 1
            continue

        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            properties = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), PALE_BLUE)
            properties.append(shading)
            run = paragraph.add_run(" ".join(quote_lines))
            run.italic = True
            run.font.color.rgb = INK
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and not is_special_line(lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))


def main() -> None:
    for required in (SOURCE, TEMPLATE, ARCHITECTURE_IMAGE, SEQUENCE_IMAGE):
        if not required.exists():
            raise SystemExit(f"Missing required input: {required}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    working_copy = OUTPUT.with_suffix(".working.docx")
    shutil.copy2(TEMPLATE, working_copy)
    document = Document(working_copy)
    clear_body(document)
    configure_document(document)
    add_cover(document)
    render_body(document, SOURCE.read_text(encoding="utf-8"))
    document.save(OUTPUT)
    working_copy.unlink(missing_ok=True)
    print(OUTPUT)


if __name__ == "__main__":
    main()
